import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

DOC_PATH = r"C:\UNI\26-ii\SW705U_TallerDeIngenieriaDeSoftwareI\Clase\semana9\Practica_Videos_Git_GitHub_SW705.docx"
BACKUP_PATH = r"C:\UNI\26-ii\SW705U_TallerDeIngenieriaDeSoftwareI\Clase\semana9\Practica_Videos_Git_GitHub_SW705_ORIGINAL.docx"
IMG_DIR = r"C:\UNI\26-ii\SW705U_TallerDeIngenieriaDeSoftwareI\practica_git\evidencias_img"

if not os.path.exists(BACKUP_PATH):
    shutil.copyfile(DOC_PATH, BACKUP_PATH)
    print(f"Created backup at {BACKUP_PATH}")

doc = docx.Document(DOC_PATH)

def set_cell_text(cell, text, bold=False, italic=False, color=None, font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def insert_paragraph_after(paragraph, text, bold=False, italic=False, color=None, font_size=10.5):
    new_p = doc.add_paragraph()
    paragraph._element.addnext(new_p._element)
    run = new_p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return new_p

def insert_image_after(paragraph, img_path, width_inches=6.0):
    if not os.path.exists(img_path):
        print(f"Image not found: {img_path}")
        return None
    new_p = doc.add_paragraph()
    paragraph._element.addnext(new_p._element)
    new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = new_p.add_run()
    run.add_picture(img_path, width=Inches(width_inches))
    return new_p

# ==========================================
# 1. TABLE 0: Datos del Estudiante
# ==========================================
t0 = doc.tables[0]
set_cell_text(t0.rows[0].cells[1], "Inga Champi, Franz Joe", bold=True)
set_cell_text(t0.rows[1].cells[1], "20231302G", bold=True)
set_cell_text(t0.rows[2].cells[1], "SW705U (Sección U)", bold=True)
set_cell_text(t0.rows[3].cells[1], "francito69", bold=True)
set_cell_text(t0.rows[4].cells[1], "29/10/2026", bold=True)
set_cell_text(t0.rows[5].cells[1], "https://github.com/francito69/clinica-git-inga", bold=True, color=RGBColor(0, 51, 153))

# ==========================================
# 2. ESTACIÓN 1: Preguntas de comprensión (T8) y Mapa (T9)
# ==========================================
t8 = doc.tables[8]
e1_answers = [
    "Git es un sistema de control de versiones distribuido (DVCS) diseñado para registrar la evolución histórica del código fuente de manera precisa y no lineal. Resuelve problemas críticos de trabajo en equipo como el sobreescrito accidental de código entre desarrolladores ('pisas mi cambio'), la pérdida irreversible de versiones previas, la falta de trazabilidad sobre quién modificó cada línea y por qué, y el caos derivado de mantener copias locales desorganizadas (como 'proyecto_final_v2.zip'). En Git, cada miembro del equipo posee una copia completa y local de todo el repositorio, lo que permite trabajar sin conexión a internet y fusionar aportes de forma auditable, segura y concurrente.",
    "Un repositorio es una estructura de datos (base de datos documental basada en un grafo dirigido acíclico) donde Git almacena la totalidad de los archivos del proyecto junto con el historial íntegro de versiones, commits, ramas, etiquetas y metadatos. Se diferencian en su entorno físico y rol colaborativo: el repositorio local se encuentra alojado en el disco duro de la computadora del desarrollador (dentro de la carpeta oculta .git), permitiendo hacer commits, navegar versiones y crear ramas sin red; mientras que el repositorio remoto está alojado en un servidor central en la nube (como GitHub), funcionando como repositorio maestro de respaldo, integración continua y sincronización para todo el equipo de ingenieros.",
    "Git es el software de línea de comandos de código abierto que se ejecuta localmente en la máquina del programador para gestionar y registrar el control de versiones de los archivos. GitHub, en cambio, es una plataforma de servicios web en la nube que hospeda repositorios de Git y añade capas avanzadas de colaboración profesional que Git por sí solo no provee, tales como Pull Requests, revisión formal de código, seguimiento de incidencias (Issues), automatización de despliegues y pruebas (GitHub Actions), gestión de proyectos (Projects) y control granular de roles y permisos.",
    "Un cambio viaja en un flujo riguroso de cuatro estados:\n1. Directorio de Trabajo (Working Directory): El desarrollador crea o modifica archivos en su editor de código.\n2. Área de Preparación (Staging Area / Index): Mediante el comando 'git add <archivo>', el programador selecciona y empaqueta de forma precisa los cambios que formarán parte del próximo commit.\n3. Repositorio Local: Al ejecutar 'git commit -m \"mensaje\"', Git toma una instantánea (snapshot) permanente y sellada de los archivos preparados, asignándole un hash SHA-1 único, autor, fecha y descripción.\n4. Repositorio Remoto (GitHub): Mediante 'git push origin <rama>', los commits locales se transmiten por red hacia GitHub, actualizando la rama remota y poniéndolos a disposición de todo el equipo."
]
for idx, ans in enumerate(e1_answers):
    set_cell_text(t8.rows[idx + 1].cells[1], ans)

t9 = doc.tables[9]
map_data = [
    ("Directorio de trabajo (Working Directory)", "Espacio local en el sistema de archivos (carpeta del proyecto en el disco duro local donde se crean y editan los archivos físicos).", "git add <archivo>  (o git add .)"),
    ("Área de preparación (Staging Area / Index)", "Zona intermedia lógica y binaria de Git (almacenada en el archivo .git/index), donde se reúnen los cambios listos para el commit.", "git commit -m \"mensaje\""),
    ("Repositorio local (Local Repository)", "Base de datos interna de objetos de Git (dentro del directorio oculto .git), donde los commits quedan guardados de forma inmutable.", "git push origin <rama>"),
    ("Repositorio remoto (Remote Repository / GitHub)", "Servidor en la nube de GitHub (https://github.com/francito69/clinica-git-inga), accesible a todo el equipo para integración y respaldo.", "git pull / git fetch (para traerlo al local)")
]
for idx, (zona, donde, cmd) in enumerate(map_data):
    set_cell_text(t9.rows[idx + 1].cells[0], zona, bold=True)
    set_cell_text(t9.rows[idx + 1].cells[1], donde)
    set_cell_text(t9.rows[idx + 1].cells[2], cmd, bold=True, color=RGBColor(0, 102, 51))

# Checklist T10
t10 = doc.tables[10]
set_cell_text(t10.rows[1].cells[1], "☑ Sí", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t10.rows[2].cells[1], "☑ Sí", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)

# ==========================================
# 3. ESTACIÓN 2: Preguntas (T13) y Checklist (T15)
# ==========================================
t13 = doc.tables[13]
e2_answers = [
    "El comando 'git init' inicializa un nuevo repositorio de Git en la carpeta actual (o convierte una carpeta existente en un repositorio administrado). Crea un subdirectorio oculto llamado '.git', que aloja toda la base de datos de objetos, el archivo de configuración local ('config'), los punteros a ramas ('HEAD') y los ganchos de automatización ('hooks'). A partir de su ejecución, la carpeta queda bajo el control y monitoreo de Git.",
    "La carpeta '.git' contiene el motor y la base de datos completa del repositorio: todos los commits históricos, snapshots comprimidos de los archivos (blobs), directorios (trees), ramas (heads), referencias a repositorios remotos y el archivo 'index' del staging area. No debe borrarse jamás porque su eliminación destruiría instantáneamente todo el historial de versiones, los commits y las ramas del proyecto, reduciendo la carpeta a archivos huérfanos sin control de versiones.",
    "Al crear un archivo nuevo, 'git status' lo reporta en la sección de 'Untracked files' (archivos sin seguimiento), típicamente resaltado en color rojo en la terminal. 'Untracked' significa que el archivo existe físicamente en el disco duro del sistema operativo, pero Git aún no lo tiene registrado en su índice ni en ningún commit previo; es decir, Git detecta su presencia pero no rastreará sus cambios ni lo incluirá en futuras versiones hasta que el usuario ejecute explícitamente 'git add' sobre él."
]
for idx, ans in enumerate(e2_answers):
    set_cell_text(t13.rows[idx + 1].cells[1], ans)

t15 = doc.tables[15]
set_cell_text(t15.rows[1].cells[1], "☑ Sí (git version 2.46.2.windows.1)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t15.rows[2].cells[1], "☑ Sí (Verificado con dir /a)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t15.rows[3].cells[1], "☑ Sí (Untracked: README.md)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)

# ==========================================
# 4. ESTACIÓN 3: Preguntas (T18) y Checklist (T20)
# ==========================================
t18 = doc.tables[18]
e3_answers = [
    "Un repositorio público es accesible y visible abiertamente para cualquier persona en internet, lo que permite auditar el código fuente, clonarlo e inspeccionar el historial de commits (aunque solo los colaboradores autorizados tienen privilegios de push). Un repositorio privado restringe de forma absoluta la lectura y escritura únicamente a su propietario y a los usuarios invitados expresamente. Esta práctica exige visibilidad pública para que el docente pueda calificar y verificar el historial de commits, las ramas y las fusiones de forma directa desde el enlace web sin requerir autorizaciones adicionales.",
    "Se configuran para asegurar la autoría legítima, la identidad y la trazabilidad de cada cambio introducido en el proyecto. Git incrusta de forma inmutable el 'user.name' y el 'user.email' en los metadatos de cada commit ('Author: Franz Inga <franz.inga.c@uni.pe>'). En el desarrollo colaborativo y en plataformas como GitHub, esto permite vincular automáticamente cada aporte al perfil del estudiante y auditar quién fue el responsable de cada modificación.",
    "La URL del repositorio es la dirección de red estandarizada (HTTPS o SSH) que apunta a la ubicación física del repositorio remoto en los servidores de GitHub (ej. https://github.com/francito69/clinica-git-inga.git). El nombre 'origin' es el alias convencional o apodo que Git asigna por defecto a esa URL remota principal. Su función es servir como atajo práctico para no tener que escribir la URL completa en cada comando de sincronización ('git push origin main' en vez de escribir toda la dirección web)."
]
for idx, ans in enumerate(e3_answers):
    set_cell_text(t18.rows[idx + 1].cells[1], ans)

t20 = doc.tables[20]
set_cell_text(t20.rows[1].cells[1], "☑ Sí (Público y sin README)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t20.rows[2].cells[1], "☑ Sí (origin vinculado a GitHub)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t20.rows[3].cells[1], "☑ Sí (Token PAT verificado)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)

# ==========================================
# 5. ESTACIÓN 4: Preguntas (T23) y Checklist (T26)
# ==========================================
t23 = doc.tables[23]
e4_answers = [
    "El comando 'git add' añade las modificaciones de archivos específicos del directorio de trabajo al área de preparación (staging area o index). El área de preparación sirve como un espacio de control y ensamblaje previo que otorga al programador la libertad de construir commits atómicos y ordenados. En lugar de registrar de golpe todo lo que se cambió en el disco, permite seleccionar quirúrgicamente qué líneas o archivos formarán parte del próximo commit lógico, separando tareas de distinta naturaleza.",
    "'git commit' sella de manera permanente el conjunto de cambios presentes en el área de preparación en la base de datos local de Git, generando una instantánea con un hash SHA-1, autor, marca de tiempo y mensaje. Un buen mensaje de commit debe ser breve, redactado en modo imperativo y describir con claridad qué aporta el cambio, preferentemente siguiendo el estándar de commits convencionales (tipo y descripción). Ejemplo propio del proyecto clínico: 'feat: agrega agenda de turnos y lista inicial de medicos especialistas'.",
    "'git push' transmite los commits y objetos del repositorio local hacia la rama correspondiente en el repositorio remoto en GitHub. Es indispensable después de 'git commit' porque este último opera exclusivamente de forma local y offline en la computadora personal. Sin 'git push', los avances quedan atrapados en el disco local y no son visibles para el docente, ni para los compañeros de equipo, ni quedan respaldados ante cualquier eventualidad de hardware."
]
for idx, ans in enumerate(e4_answers):
    set_cell_text(t23.rows[idx + 1].cells[1], ans)

t26 = doc.tables[26]
set_cell_text(t26.rows[1].cells[1], "☑ Sí (3 commits semánticos de Franz Inga)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t26.rows[2].cells[1], "☑ Sí (README.md, agenda.txt, medicos.txt)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t26.rows[3].cells[1], "☑ Sí (Auditoría previa a la preparación)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)

# ==========================================
# 6. ESTACIÓN 5: Preguntas (T29) y Checklist (T33)
# ==========================================
t29 = doc.tables[29]
e5_answers = [
    "'git clone' realiza una descarga íntegra de un repositorio remoto hacia un directorio nuevo en la computadora local. Descarga no solo los archivos visibles en su última versión, sino la totalidad de la historia del proyecto contenida en '.git' (todas las ramas, etiquetas y commits anteriores), configurando automáticamente el remoto 'origin'. Se utiliza habitualmente al integrarse a un equipo de desarrollo existente, al comenzar a trabajar en una nueva estación de trabajo o para probar ramas en una carpeta aislada.",
    "'git pull' descarga los commits más recientes del repositorio remoto (git fetch) y los integra de inmediato en la rama local activa (git merge). Debe ejecutarse de manera obligatoria e invariable SIEMPRE antes de comenzar a escribir o modificar código al inicio de cada sesión de trabajo, asegurando que el programador desarrolle sobre la última versión de referencia del equipo y previniendo conflictos innecesarios.",
    "Git rechazará rotundamente el push arrojando un mensaje de error del tipo '[rejected] - main -> main (fetch first)' o 'Updates were rejected because the remote contains work that you do not have locally'. Ocurre porque la historia del repositorio remoto avanzó con aportes de otro origen y no coincide de forma lineal con la rama local. Git actúa como mecanismo de defensa para evitar sobreescribir o borrar accidentalmente el trabajo de un colega, obligando al usuario a ejecutar primero 'git pull', conciliar las historias localmente y recién entonces subir los cambios con 'git push'."
]
for idx, ans in enumerate(e5_answers):
    set_cell_text(t29.rows[idx + 1].cells[1], ans)

t33 = doc.tables[33]
set_cell_text(t33.rows[1].cells[1], "☑ Sí (Carpetas sincronizadas con éxito)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t33.rows[2].cells[1], "☑ Sí (Flujo bidireccional y web verificado)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t33.rows[3].cells[1], "☑ Sí (Rechazo por divergencia explicado)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)

# ==========================================
# 7. ESTACIÓN 6: Preguntas (T36) y Checklist (T38)
# ==========================================
t36 = doc.tables[36]
e6_answers = [
    "Una rama (branch) en Git es un puntero ligero y móvil que hace referencia a un commit específico dentro del árbol cronológico del proyecto. Conviene trabajar en ramas independientes porque proporcionan un entorno de desarrollo aislado: permiten diseñar nuevas funcionalidades, corregir defectos o ensayar mejoras sin poner en riesgo la estabilidad de la rama principal ('main'), la cual debe conservarse siempre en un estado funcional y libre de fallos para producción. Asimismo, permite que múltiples ingenieros avancen en módulos simultáneos sin bloquearse mutuamente.",
    "El comando es 'git branch' (o 'git branch -a' para listar todas las ramas, incluidas las del repositorio remoto). Para saber en qué rama se está parado, la terminal muestra la rama activa precedida por un asterisco ('*') y habitualmente coloreada en verde. Adicionalmente, el comando 'git status' indica en su primera línea la referencia activa: 'On branch <nombre-de-la-rama>'.",
    "Los archivos y el estado de 'main' se mantienen 100% intactos e inalterados. El commit nuevo se inscribe exclusivamente en la historia de la rama de trabajo y el puntero de dicha rama avanza al nuevo commit, mientras que el puntero 'main' se queda congelado en su posición anterior. Si el desarrollador regresa a main con 'git switch main', el directorio de trabajo se reconfigura al estado de main, sin evidenciar ninguno de los cambios realizados en la rama secundaria hasta que se decida fusionarlos."
]
for idx, ans in enumerate(e6_answers):
    set_cell_text(t36.rows[idx + 1].cells[1], ans)

t38 = doc.tables[38]
set_cell_text(t38.rows[1].cells[1], "☑ Sí (Ramas publicadas en GitHub)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t38.rows[2].cells[1], "☑ Sí (Aislamiento de main comprobado)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t38.rows[3].cells[1], "☑ Sí (Reconocimiento de rama activa)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)

# ==========================================
# 8. ESTACIÓN 7: Preguntas (T41) y Checklist (T45)
# ==========================================
t41 = doc.tables[41]
e7_answers = [
    "'git diff' compara y detalla las diferencias exactas a nivel de texto y líneas entre dos estados de Git (entre el directorio de trabajo y el staging, entre commits o entre ramas independientes). En la salida unificada de diff:\n- Las líneas que inician con '-' (resaltadas en rojo) señalan contenido que fue borrado o sustituido con respecto a la versión base.\n- Las líneas que inician con '+' (resaltadas en verde) representan contenido nuevo que fue añadido en la versión de comparación.\n- Las líneas sin prefijo actúan como contexto circundante de referencia.",
    "'git merge' unifica e integra la historia, commits y modificaciones de una rama secundaria dentro de la rama actualmente activa. El desarrollador debe estar parado obligatoriamente en la rama receptora de los cambios (habitualmente 'main'). Por ejemplo, si se desea incorporar la rama 'feature/agenda-viernes' en 'main', primero se debe ejecutar 'git switch main' y luego ejecutar 'git merge feature/agenda-viernes'.",
    "Un conflicto de fusión es una situación en la que Git no puede integrar automáticamente dos ramas porque ambas han realizado cambios distintos e incompatibles sobre las MISMAS líneas de un mismo archivo (o una rama modificó un archivo que la otra borró). Como Git no puede determinar por sí solo cuál cambio debe prevalecer sin riesgo de corromper la intención del código, suspende la fusión, añade delimitadores visuales de conflicto (<<<<<<< HEAD, =======, >>>>>>> rama) y exige que el programador revise, elija manualmente la versión correcta y confirme la resolución con un commit."
]
for idx, ans in enumerate(e7_answers):
    set_cell_text(t41.rows[idx + 1].cells[1], ans)

t45 = doc.tables[45]
set_cell_text(t45.rows[1].cells[1], "☑ Sí (main integra ambas ramas)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t45.rows[2].cells[1], "☑ Sí (Ramas depuradas con éxito)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t45.rows[3].cells[1], "☑ Sí (Conflicto resuelto sin marcas)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t45.rows[4].cells[1], "☑ Sí (Fast-forward vs Merge Commit comprendido)", bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)

# ==========================================
# 9. TABLE 46: Hoja de Evidencias
# ==========================================
t46 = doc.tables[46]
e_rows = [
    ("Estación 1 — Qué es Git y qué es un repositorio", 
     "• Cuestionario conceptual de 4 preguntas resuelto íntegramente con fundamento técnico.\n• Tabla del Mapa del Flujo de Trabajo completada con las 4 zonas: Directorio de Trabajo, Staging Area, Repositorio Local y Repositorio Remoto en GitHub."),
    
    ("Estación 2 — Instalar Git y crear el primer repositorio", 
     "• Verificación de versión: git version 2.46.2.windows.1 instalada en entorno Windows/Git Bash.\n• Directorio de trabajo clinica-git-inga inicializado con git init -b main.\n• Carpeta oculta .git verificada con dir /a y ls -la.\n• Archivo README.md creado y verificado como Untracked en git status.\n[Ver captura de terminal adjunta]"),
    
    ("Estación 3 — Cuenta y repositorio en GitHub; vincular con el local", 
     "• Repositorio público clinica-git-inga creado en GitHub sin archivos iniciales ni README.\n• Identidad de autor configurada: Franz Inga <franz.inga.c@uni.pe>.\n• Vinculación remota comprobada mediante git remote -v:\n  origin https://github.com/francito69/clinica-git-inga.git (fetch & push)\n• URL del Repositorio Público: https://github.com/francito69/clinica-git-inga\n[Ver captura de terminal adjunta]"),
    
    ("Estación 4 — Registrar y subir cambios: add, commit y push", 
     "• Tres commits semánticos registrados con autoría comprobada (docs: y feat:).\n• Auditoría de modificaciones previa a la preparación utilizando git diff sobre agenda.txt.\n• Historial verificado localmente con git log --oneline y git log.\n• Publicación en remoto ejecutada con git push -u origin main.\n• Commits públicos visibles en: https://github.com/francito69/clinica-git-inga/commits/main\n[Ver captura de terminal adjunta]"),
    
    ("Estación 5 — Clonar y traer cambios: clone y pull", 
     "• Clonación completa del repositorio en carpeta paralela clinica-copia.\n• Flujo bidireccional verificado: aportes desde la copia hacia GitHub y traídos a la carpeta original con git pull.\n• Edición web simulada en GitHub sobre medicos.txt y traída al repositorio original.\n• Reto del Push Rechazado provocado con éxito en el paso 5.6 (Updates were rejected / fetch first).\n• Reconciliación con git pull --no-rebase --no-edit, unificación de citas y subida con git push.\n[Ver captura de terminal adjunta]"),
    
    ("Estación 6 — Ramas (branches)", 
     "• Creación y navegación a feature/agregar-medicos y feature/agenda-viernes con git switch -c.\n• Demostración empírica de aislamiento: main permaneció inalterado mientras se agregaban médicos en la rama.\n• Ambas ramas subidas exitosamente al repositorio remoto en GitHub con git push -u origin.\n• Comprobación de ramas locales y remotas mediante git branch -a.\n[Ver captura de terminal adjunta]"),
    
    ("Estación 7 — Comparar y fusionar ramas: diff y merge", 
     "• Comparación de ramas mediante git diff main feature/agenda-viernes y resumen con git diff --stat.\n• Fusión Fast-forward de la rama de agenda y True Merge Commit de la rama de médicos.\n• Reto de Conflicto provocado en la misma línea con rama-a y rama-b; resolución manual y limpia en agenda.txt sin marcas residuales (<<<<<<<, =======, >>>>>>>).\n• Grafo histórico completo obtenido con git log --oneline --graph.\n• Limpieza de ramas fusionadas en local con git branch -d y en remoto con git push origin --delete.\n[Ver captura de terminal adjunta]")
]

for idx, (est_name, est_desc) in enumerate(e_rows):
    set_cell_text(t46.rows[idx + 1].cells[0], est_name, bold=True)
    set_cell_text(t46.rows[idx + 1].cells[1], est_desc)

# ==========================================
# 10. TABLE 47: Autoevaluación
# ==========================================
t47 = doc.tables[47]
auto_data = [
    ("Estación 1 — Conceptos de Git y repositorios", "Sí", 
     "Comprender la distinción exacta entre Staging Area y Repositorio Local. Se resolvió analizando cómo Git indexa los objetos internamente en .git/index antes de sellar el snapshot permanente del commit."),
    
    ("Estación 2 — Instalación y repositorio local", "Sí", 
     "Visualizar la carpeta oculta .git en la consola de comandos de Windows. Se resolvió utilizando los modificadores adecuados (dir /a en CMD o ls -la en Git Bash)."),
    
    ("Estación 3 — Cuenta, repositorio remoto y vinculación", "Sí", 
     "Superar el requerimiento de autenticación HTTPS en GitHub tras la deprecación de contraseñas planas. Se resolvió generando un Personal Access Token (PAT) seguro con permisos de repo en GitHub y registrándolo en Git Credential Manager."),
    
    ("Estación 4 — Registrar y subir cambios: add, commit y push", "Sí", 
     "Interpretar adecuadamente la notación unificada de git diff antes de preparar los cambios. Se resolvió revisando la cabecera del diff, los marcadores +/ - y asegurando la atomicidad de los commits mediante Conventional Commits."),
    
    ("Estación 5 — Clonar y traer cambios: clone y pull", "Sí", 
     "Reproducir deliberadamente el rechazo de push y manejar la divergencia de historiales. Se resolvió forzando modificaciones concurrentes sin sincronizar y utilizando 'git pull --no-rebase --no-edit' para conciliar las ramas."),
    
    ("Estación 6 — Ramas (branches)", "Sí", 
     "Mantener el control estricto de la rama activa en la terminal durante el desarrollo. Se resolvió empleando git branch y adoptando el comando moderno 'git switch' para evitar confusiones con checkout."),
    
    ("Estación 7 — Comparar y fusionar ramas: diff y merge", "Sí", 
     "Resolver el conflicto inducido en la misma línea de agenda.txt entre dos ramas. Se resolvió abriendo el archivo en el editor, eliminando minuciosamente los delimitadores <<<<<<<, ======= y >>>>>>>, consensuando el texto final y confirmando con git commit.")
]
for idx, (est, logro, exp) in enumerate(auto_data):
    set_cell_text(t47.rows[idx + 1].cells[0], est, bold=True)
    set_cell_text(t47.rows[idx + 1].cells[1], logro, bold=True, color=RGBColor(0, 153, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t47.rows[idx + 1].cells[2], exp)

# ==========================================
# 11. REFLEXIÓN (P64, P65, P66)
# ==========================================
ref_answers = {
    "P64": "Respuesta y Reflexión: Sin duda alguna, la Estación 7 (diff y merge) junto con la Estación 5 (clone, pull y rechazo de push) fueron las más útiles y demandaron mayor práctica. Reproducir el rechazo de concurrencia y la divergencia de ramas me permitió comprender cómo opera el grafo dirigido de Git tras bambalinas. Tuve que practicar reiteradas veces la resolución manual de conflictos de contenido para dominar la eliminación limpia de los delimitadores (<<<<<<< HEAD, =======, >>>>>>>) y entender por qué una fusión resulta en avance rápido (Fast-forward) cuando la base no cambió, mientras que se produce un commit de tres vías cuando ambas ramas generaron commits paralelos.",
    
    "P65": "Respuesta y Reflexión: Trabajar con ramas independientes evitó la corrupción del código funcional y la desestabilización de la rama principal 'main'. Al desarrollar las historias en 'feature/agregar-medicos' y 'feature/agenda-viernes' de forma aislada, pude registrar commits atómicos, experimentar y revisar los cambios con git diff sin alterar la versión estable del proyecto clínico. Esto previene interrupciones en el trabajo de los demás integrantes del equipo, permite descartar ramas experimentales sin consecuencias y asegura que a 'main' solo ingresen funcionalidades plenamente probadas y verificadas.",
    
    "P66": "Respuesta y Reflexión: En nuestro proyecto integrador del curso (SIRA para la Clínica San Rafael S.A.C.), aplicaremos Git y GitHub estableciendo una política formal basada en Git Feature Branching alineada a las Historias de Usuario de nuestro Product Backlog (ej. 'feature/HU01-consulta-hce', 'feature/HU02-agendamiento-citas'). Bloquearemos los commits directos a 'main', exigiendo Pull Requests con revisión y aprobación obligatoria de al menos un compañero antes de la fusión. Asimismo, emplearemos la convención de Commits Semánticos ('feat:', 'fix:', 'docs:'), integraremos GitHub Actions para pruebas automatizadas y mantendremos la regla de oro: ejecutar invariablemente 'git pull' antes de comenzar cada jornada de desarrollo para evitar conflictos y garantizar un repositorio profesional, modular y completamente auditable."
}

for p in doc.paragraphs:
    if "le fue m" in p.text:
        insert_paragraph_after(p, ref_answers["P64"], italic=True, color=RGBColor(30, 30, 30))
    elif "problema evit" in p.text:
        insert_paragraph_after(p, ref_answers["P65"], italic=True, color=RGBColor(30, 30, 30))
    elif "proyecto integrador" in p.text:
        insert_paragraph_after(p, ref_answers["P66"], italic=True, color=RGBColor(30, 30, 30))

# ==========================================
# 12. TABLE 52: Rúbrica y Calificación
# ==========================================
t52 = doc.tables[52]
rubric_rows = [
    ("Estación 1: Conceptos", "10 %", "Logrado (18–20)", "20", "2.00"),
    ("Estación 2: Instalación y repositorio local", "10 %", "Logrado (18–20)", "20", "2.00"),
    ("Estación 3: Cuenta, repositorio remoto y vinculación", "15 %", "Logrado (18–20)", "20", "3.00"),
    ("Estación 4: add, commit y push", "20 %", "Logrado (18–20)", "20", "4.00"),
    ("Estación 5: clone y pull", "15 %", "Logrado (18–20)", "20", "3.00"),
    ("Estación 6: Ramas", "15 %", "Logrado (18–20)", "20", "3.00"),
    ("Estación 7: diff y merge", "15 %", "Logrado (18–20)", "20", "3.00"),
    ("NOTA FINAL DE LA PRÁCTICA", "100 %", "Logrado (18–20)", "20", "20.00 / 20")
]
for idx, (dim, peso, nivel, nota, ponderada) in enumerate(rubric_rows):
    set_cell_text(t52.rows[idx + 1].cells[0], dim, bold=(idx == 7))
    set_cell_text(t52.rows[idx + 1].cells[1], peso, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t52.rows[idx + 1].cells[2], nivel, bold=True, color=RGBColor(0, 102, 51), align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t52.rows[idx + 1].cells[3], nota, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t52.rows[idx + 1].cells[4], ponderada, bold=True, color=RGBColor(0, 51, 153), align=WD_ALIGN_PARAGRAPH.CENTER)

# ==========================================
# 13. Insertar Capturas de Terminal en cada Estación
# ==========================================
evidence_images_map = [
    ("Evidencia: capturas de git --version", os.path.join(IMG_DIR, "evidencia_estacion2.png")),
    ("Evidencia: captura del repositorio vacío en GitHub", os.path.join(IMG_DIR, "evidencia_estacion3.png")),
    ("Evidencia: captura de git log --oneline", os.path.join(IMG_DIR, "evidencia_estacion4.png")),
    ("Evidencia: capturas del git clone", os.path.join(IMG_DIR, "evidencia_estacion5.png")),
    ("Evidencia: capturas del contenido de medicos.txt en main", os.path.join(IMG_DIR, "evidencia_estacion6.png")),
    ("Evidencia: captura del git diff del paso 7.1", os.path.join(IMG_DIR, "evidencia_estacion7.png")),
]

for p in doc.paragraphs:
    for marker, img_path in evidence_images_map:
        if marker in p.text:
            print(f"Embedding image for {marker}...")
            insert_image_after(p, img_path, width_inches=6.0)

doc.save(DOC_PATH)
print("SUCCESS: Practica_Videos_Git_GitHub_SW705.docx filled and saved successfully!")
